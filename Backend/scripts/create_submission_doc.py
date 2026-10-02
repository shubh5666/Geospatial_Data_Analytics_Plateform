"""Build the Darukaa.Earth application submission document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "Darukaa_Earth_Submission.docx"
REPOSITORY_URL = "https://github.com/shubh5666/Geospatial_Data_Analytics_Plateform"
FRONTEND_URL = "https://geospatial-data-analytics-plateform-frontend-9tigroc7g.vercel.app/"
BACKEND_URL = "https://geospatial-data-analytics-plateform.onrender.com"


def set_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def format_paragraph(paragraph, after=6, before=0, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, after=8 if level == 1 else 6, before=16 if level == 1 else 12)
    run = paragraph.add_run(text)
    set_font(run, 16 if level == 1 else 13, True, "2E74B5")
    return paragraph


def bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    format_paragraph(paragraph, after=4, line=1.167)
    set_font(paragraph.add_run(text))
    return paragraph


def add_label_value(doc, label, value=None, url=None):
    p = doc.add_paragraph()
    format_paragraph(p, after=4)
    set_font(p.add_run(f"{label}: "), bold=True, color="1F4D78")
    if url:
        add_hyperlink(p, value, url)
    else:
        set_font(p.add_run(value))
    return p


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_paragraph(title, after=4)
    set_font(title.add_run("Darukaa.Earth"), size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    format_paragraph(subtitle, after=14)
    set_font(subtitle.add_run("Full-Stack Developer Hackathon Submission"), size=13, color="4F5B66")

    meta = doc.add_table(rows=2, cols=2)
    table_geometry(meta, [Inches(1.3), Inches(5.2)])
    values = [("Repository", REPOSITORY_URL), ("Stack", "React, FastAPI, PostgreSQL/PostGIS, JWT, GitHub Actions")]
    for row, (label, value) in zip(meta.rows, values):
        shade(row.cells[0], "F2F4F7")
        p0 = row.cells[0].paragraphs[0]
        format_paragraph(p0, after=0)
        set_font(p0.add_run(label), bold=True, color="1F4D78")
        p1 = row.cells[1].paragraphs[0]
        format_paragraph(p1, after=0)
        if label == "Repository":
            add_hyperlink(p1, value, value)
        else:
            set_font(p1.add_run(value))

    heading(doc, "Submission links")
    add_label_value(doc, "GitHub repository", REPOSITORY_URL, REPOSITORY_URL)
    add_label_value(doc, "Frontend live demo", FRONTEND_URL, FRONTEND_URL)
    add_label_value(doc, "Backend API", BACKEND_URL, BACKEND_URL)
    add_label_value(doc, "API documentation", f"{BACKEND_URL}/docs", f"{BACKEND_URL}/docs")

    heading(doc, "Architecture overview")
    p = doc.add_paragraph()
    format_paragraph(p)
    set_font(p.add_run("Darukaa.Earth is a full-stack geospatial analytics dashboard for carbon and biodiversity projects. The React/Vite frontend presents an authenticated workspace with project and site management, polygon drawing, map visualization, analytics charts, GeoJSON export, and a synthetic demo workspace. A FastAPI backend exposes JWT-protected REST endpoints and enforces ownership checks for every project and site query. PostgreSQL with PostGIS stores validated site polygons and associated measurements."))

    heading(doc, "Database schema")
    for item in [
        "users - UUID account records, email, full name, and Argon2 password hash.",
        "projects - user-owned project records with name and description.",
        "sites - project-owned polygon boundaries stored as geometry(POLYGON, 4326) in PostGIS.",
        "site_measurements - dated carbon and biodiversity metrics for each site.",
    ]:
        bullet(doc, item)

    heading(doc, "Local setup")
    for item in [
        "Configure Backend/.env from Backend/.env.example with PostgreSQL/PostGIS and JWT settings.",
        "Install backend packages, run migrations with python -m app.init_db, then start FastAPI with uvicorn app.main:app --app-dir Backend --reload.",
        "Install workspace dependencies with npm ci, then run npm run dev for the React frontend.",
        "Set VITE_MAPBOX_ACCESS_TOKEN in Frontend/.env when Mapbox basemaps are required; the coordinate-grid map works without it.",
    ]:
        bullet(doc, item)

    heading(doc, "CI/CD and code quality")
    for item in [
        "GitHub Actions runs backend Ruff linting and PostGIS integration tests, plus frontend linting, unit tests, and production builds on pushes and pull requests.",
        "Husky and lint-staged provide pre-commit formatting and linting checks; Prettier enforces consistent formatting.",
        "render.yaml defines the Render FastAPI/PostGIS deployment. vercel.json defines the Vercel frontend build and output directory. Both platforms deploy from the main branch after configuration.",
    ]:
        bullet(doc, item)

    heading(doc, "Reviewer notes")
    p = doc.add_paragraph()
    format_paragraph(p)
    set_font(p.add_run("The repository README contains full local-run, schema, analytics, security, testing, and deployment instructions. The public Vercel frontend and Render API URLs are included above. If the repository is made private, access must be granted to the Darukaa hiring-team email accounts listed in the assignment."))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(footer, after=0)
    set_font(footer.add_run("Darukaa.Earth submission"), size=9, color="6B7280")

    doc.core_properties.title = "Darukaa.Earth Hackathon Submission"
    doc.core_properties.subject = "Full-stack developer hackathon submission"
    doc.core_properties.author = "Shubh"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
